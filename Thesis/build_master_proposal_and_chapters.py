import os
import re
import shutil
import markdown
import docx
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

thesis_dir = "/home/radhi/Documents/AUV Development/Thesis"
docx_dir = os.path.join(thesis_dir, "docx_export")
html_dir = os.path.join(thesis_dir, "html_export")
github_thesis_dir = "/home/radhi/Documents/AUV_GitHub_Upload/Thesis"
github_docx_dir = os.path.join(github_thesis_dir, "docx_export")
github_html_dir = os.path.join(github_thesis_dir, "html_export")

os.makedirs(docx_dir, exist_ok=True)
os.makedirs(html_dir, exist_ok=True)
os.makedirs(github_docx_dir, exist_ok=True)
os.makedirs(github_html_dir, exist_ok=True)

logo_path = "/home/radhi/Documents/AUV Development/unhas_logo.png"

# Printable width for B5: 176 mm - 45 mm (margins) = 131 mm
PRINTABLE_WIDTH = Mm(131)

OFFICIAL_TITLE_MD = "ANALISIS KINEMATIKA, DINAMIKA, DAN ESTIMASI KEADAAN OPTIMAL *KALMAN FILTER* UNTUK *VISION-BASED TRACKING* PADA *OVER-ACTUATED 8-THRUSTER 6-DOF VECTORED AUV*"
OFFICIAL_SUBTITLE_EN = "(Analysis of Kinematics, Dynamics, and Optimal Kalman Filter State Estimation for Vision-Based Tracking on an Over-Actuated 8-Thruster 6-DOF Vectored AUV)"

def setup_unhas_section(doc, is_front_matter=False, start_page=1, add_page_number=True):
    sec = doc.sections[0]
    sec.page_width = Mm(176)
    sec.page_height = Mm(250)
    sec.top_margin = Mm(22.5)
    sec.bottom_margin = Mm(22.5)
    sec.left_margin = Mm(22.5)
    sec.right_margin = Mm(22.5)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if add_page_number:
        sectPr = sec._sectPr
        if is_front_matter:
            sec.different_first_page_header_footer = True
            sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>'))
        else:
            sec.different_first_page_header_footer = False
            sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal" w:start="{start_page}"/>'))

        p_hdr = sec.header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.paragraph_format.space_after = Pt(0)
        p_hdr.paragraph_format.line_spacing = 1.0
        r_hdr = p_hdr.add_run()
        r_hdr.font.name = 'Arial'
        r_hdr.font.size = Pt(9.5)
        fld_xml = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_hdr._p.append(fld_xml)

def parse_inline_runs(text):
    # Matches $$...$$, **...**, *...*, `...`
    pattern = re.compile(r'(\$\$[^\$]+\$\$|\*\*[^\*]+\*\*|\*[^\*]+\*|`[^`]+`)')
    parts = pattern.split(text)
    runs = []
    for part in parts:
        if not part:
            continue
        if part.startswith('$$') and part.endswith('$$'):
            runs.append(('math', part))
        elif part.startswith('**') and part.endswith('**'):
            val = part[2:-2].strip('*#').strip()
            if val:
                runs.append(('bold', val))
        elif part.startswith('*') and part.endswith('*'):
            val = part[1:-1].strip('*#').strip()
            if val:
                runs.append(('italic', val))
        elif part.startswith('`') and part.endswith('`'):
            val = part[1:-1].strip()
            if val:
                runs.append(('code', val))
        else:
            # TEXT TOKEN:
            # Strip any unparsed, stray '*' or '#' markdown markers that are not part of LaTeX math
            clean_text = part.replace('*', '').replace('#', '')
            if clean_text:
                runs.append(('text', clean_text))
    return runs

def add_styled_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, line_spacing=1.15, first_indent=Mm(10), bold_all=False, italic_all=False, font_size=10):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if first_indent:
        p.paragraph_format.first_line_indent = first_indent

    tokens = parse_inline_runs(text)
    for t_type, t_val in tokens:
        run = p.add_run(t_val)
        run.font.name = 'Arial'
        run.font.size = Pt(font_size)
        if bold_all or t_type == 'bold':
            run.bold = True
        if italic_all or t_type == 'italic':
            run.italic = True
        if t_type == 'code':
            run.font.name = 'Consolas'
            run.font.size = Pt(font_size - 0.5)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def add_heading_1(doc, title_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    clean_title = title_text.replace('*', '').replace('#', '')
    run = p.add_run(clean_title)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.bold = True
    return p

def add_heading_2(doc, title_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    tokens = parse_inline_runs(title_text)
    for t_type, t_val in tokens:
        run = p.add_run(t_val)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.bold = True
        if t_type == 'italic':
            run.italic = True
    return p

def add_heading_3(doc, title_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    tokens = parse_inline_runs(title_text)
    for t_type, t_val in tokens:
        run = p.add_run(t_val)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.bold = True
        if t_type == 'italic':
            run.italic = True
    return p

def add_leader_line(doc, left_text, page_str, indent_mm=0, bold=False, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(PRINTABLE_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if indent_mm > 0:
        p.paragraph_format.left_indent = Mm(indent_mm)

    tokens = parse_inline_runs(left_text)
    for t_type, t_val in tokens:
        r1 = p.add_run(t_val)
        r1.font.name = 'Arial'
        r1.font.size = Pt(9.5)
        if bold or t_type == 'bold':
            r1.bold = True
        if t_type == 'italic':
            r1.italic = True

    r2 = p.add_run(f"\t{page_str}")
    r2.font.name = 'Arial'
    r2.font.size = Pt(9.5)
    if bold:
        r2.bold = True
    return p

def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(r'''
        <w:tcBorders %s >
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tcBorders>
    ''' % nsdecls('w'))
    tcPr.append(borders)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def build_front_matter(doc):
    # 2.1 Cover Page (Unnumbered)
    p_cov = doc.add_paragraph()
    p_cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov.paragraph_format.space_before = Pt(10)
    p_cov.paragraph_format.space_after = Pt(18)
    r = p_cov.add_run("PROPOSAL TUGAS AKHIR")
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(8)
    p_title.paragraph_format.line_spacing = 1.15
    for t_type, t_val in parse_inline_runs(OFFICIAL_TITLE_MD):
        r_t = p_title.add_run(t_val)
        r_t.font.name = 'Arial'
        r_t.font.size = Pt(11)
        r_t.bold = True
        if t_type == 'italic':
            r_t.italic = True

    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subtitle.paragraph_format.space_after = Pt(24)
    r_sub = p_subtitle.add_run(OFFICIAL_SUBTITLE_EN)
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(9.5)
    r_sub.italic = True

    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(28)
        p_logo.add_run().add_picture(logo_path, width=Mm(28), height=Mm(35))
    else:
        doc.add_paragraph()

    p_author_lbl = doc.add_paragraph()
    p_author_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author_lbl.paragraph_format.space_after = Pt(2)
    r_by = p_author_lbl.add_run("Disusun oleh:")
    r_by.font.name = 'Arial'
    r_by.font.size = Pt(10)

    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_after = Pt(2)
    r_name = p_author.add_run("MUH. RADHI SYAFIQ GHANIM. S")
    r_name.font.name = 'Arial'
    r_name.font.size = Pt(10)
    r_name.bold = True

    p_nim = doc.add_paragraph()
    p_nim.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nim.paragraph_format.space_after = Pt(36)
    r_nim = p_nim.add_run("NIM. D021201006")
    r_nim.font.name = 'Arial'
    r_nim.font.size = Pt(10)
    r_nim.bold = True

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(0)
    p_inst.paragraph_format.line_spacing = 1.15
    r_inst = p_inst.add_run(
        "DEPARTEMEN TEKNIK MESIN\n"
        "FAKULTAS TEKNIK\n"
        "UNIVERSITAS HASANUDDIN\n"
        "MAKASSAR\n"
        "2026"
    )
    r_inst.font.name = 'Arial'
    r_inst.font.size = Pt(10)
    r_inst.bold = True

    doc.add_page_break()

    # 2.2 LEMBAR PENGESAHAN (Halaman ii)
    add_heading_1(doc, "LEMBAR PENGESAHAN\nPROPOSAL TUGAS AKHIR")

    p_app_title = doc.add_paragraph()
    p_app_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_app_title.paragraph_format.space_after = Pt(14)
    p_app_title.paragraph_format.line_spacing = 1.15
    for t_type, t_val in parse_inline_runs(OFFICIAL_TITLE_MD):
        r_at = p_app_title.add_run(t_val)
        r_at.font.name = 'Arial'
        r_at.font.size = Pt(10.5)
        r_at.bold = True
        if t_type == 'italic':
            r_at.italic = True

    p_disusun = doc.add_paragraph()
    p_disusun.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_disusun.paragraph_format.space_after = Pt(4)
    r_d = p_disusun.add_run("Disusun dan diajukan oleh:")
    r_d.font.name = 'Arial'
    r_d.font.size = Pt(10)

    p_stud_info = doc.add_paragraph()
    p_stud_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_stud_info.paragraph_format.space_after = Pt(14)
    r_si = p_stud_info.add_run("MUH. RADHI SYAFIQ GHANIM. S\nNIM. D021201006")
    r_si.font.name = 'Arial'
    r_si.font.size = Pt(10)
    r_si.bold = True

    p_app_intro = doc.add_paragraph()
    p_app_intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_app_intro.paragraph_format.space_after = Pt(16)
    r_ai = p_app_intro.add_run("Telah diperiksa dan disetujui untuk diseminarkan pada:\nDepartemen Teknik Mesin, Fakultas Teknik, Universitas Hasanuddin")
    r_ai.font.name = 'Arial'
    r_ai.font.size = Pt(10)

    table_sup = doc.add_table(rows=1, cols=2)
    table_sup.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table_sup.rows:
        for cell in row.cells:
            cell.width = Mm(65)
    cell_l, cell_r = table_sup.rows[0].cells
    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.paragraph_format.line_spacing = 1.15
    p_l.add_run("Pembimbing Utama,\n\n\n\n\n").font.name = 'Arial'
    r_p1 = p_l.add_run("Andi Amijoyo Mochtar, S.T., M.Sc., Ph.D.\n")
    r_p1.bold = True
    r_p1.font.name = 'Arial'
    p_l.add_run("NIP. 19760216 201012 1 002").font.name = 'Arial'

    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.paragraph_format.line_spacing = 1.15
    p_r.add_run("Pembimbing Pendamping,\n\n\n\n\n").font.name = 'Arial'
    r_p2 = p_r.add_run("......................................................\n")
    r_p2.bold = True
    r_p2.font.name = 'Arial'
    p_r.add_run("NIP. ..................................................").font.name = 'Arial'

    p_mid = doc.add_paragraph()
    p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mid.paragraph_format.space_before = Pt(24)
    p_mid.paragraph_format.space_after = Pt(10)
    p_mid.paragraph_format.line_spacing = 1.15
    p_mid.add_run("Mengetahui,\nKetua Departemen Teknik Mesin Fakultas Teknik\nUniversitas Hasanuddin,\n\n\n\n").font.name = 'Arial'
    r_hod = p_mid.add_run("Dr. Muhammad Syahid, S.T., M.T.\n")
    r_hod.bold = True
    r_hod.font.name = 'Arial'
    p_mid.add_run("NIP. 19770707 200501 1 001").font.name = 'Arial'

    doc.add_page_break()

    # 2.3 PERNYATAAN KEASLIAN (Halaman iii) - Identity strictly moved to the LEFT per user request
    add_heading_1(doc, "PERNYATAAN KEASLIAN PROPOSAL TUGAS AKHIR")

    add_styled_paragraph(doc, "Yang bertanda tangan di bawah ini:", first_indent=0, space_after=6)
    t_bio = doc.add_table(rows=4, cols=3)
    t_bio.alignment = WD_TABLE_ALIGNMENT.LEFT
    bio_data = [
        ("Nama Mahasiswa", ":", "MUH. RADHI SYAFIQ GHANIM. S"),
        ("Nomor Induk Mahasiswa (NIM)", ":", "D021201006"),
        ("Program Studi", ":", "Teknik Mesin"),
        ("Departemen", ":", "Teknik Mesin, Fakultas Teknik, Universitas Hasanuddin")
    ]
    for idx, (f1, f2, f3) in enumerate(bio_data):
        row = t_bio.rows[idx]
        row.cells[0].width = Mm(55)
        row.cells[1].width = Mm(5)
        row.cells[2].width = Mm(71)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.add_run(f1).font.name = 'Arial'

        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.add_run(f2).font.name = 'Arial'

        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_val = p2.add_run(f3)
        r_val.font.name = 'Arial'
        if idx < 2:
            r_val.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_paragraph(doc, "Menyatakan dengan sesungguhnya dan penuh kesadaran bahwa Naskah Proposal Tugas Akhir yang berjudul:")
    p_prop_t = doc.add_paragraph()
    p_prop_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prop_t.paragraph_format.space_after = Pt(8)
    p_prop_t.paragraph_format.space_before = Pt(4)
    p_prop_t.add_run('"').font.name = 'Arial'
    for t_type, t_val in parse_inline_runs(OFFICIAL_TITLE_MD):
        r_pt = p_prop_t.add_run(t_val)
        r_pt.font.name = 'Arial'
        r_pt.font.size = Pt(10)
        r_pt.bold = True
        if t_type == 'italic':
            r_pt.italic = True
    p_prop_t.add_run('"').font.name = 'Arial'

    add_styled_paragraph(doc, "adalah benar-benar merupakan hasil karya ilmiah mandiri saya sendiri di bawah bimbingan komisi pembimbing yang telah ditunjuk, dan bukan merupakan karya penjiplakan (plagiasi), duplikasi, maupun saduran tanpa mencantumkan rujukan yang sah dari karya orang lain.")
    add_styled_paragraph(doc, "Apabila di kemudian hari terbukti atau dapat dibuktikan bahwa sebagian maupun keseluruhan isi naskah proposal ini mengandung unsur plagiarisme atau melanggar etika integritas akademik, maka saya bersedia menerima sanksi akademik yang tegas sesuai dengan peraturan perundang-undangan dan ketentuan hukum yang berlaku di Universitas Hasanuddin.")

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.space_before = Pt(20)
    p_sig.paragraph_format.line_spacing = 1.15
    p_sig.add_run("Makassar, 24 September 2026\nYang membuat pernyataan,\n\n\n\n\n").font.name = 'Arial'
    r_sname = p_sig.add_run("MUH. RADHI SYAFIQ GHANIM. S\n")
    r_sname.bold = True
    r_sname.font.name = 'Arial'
    p_sig.add_run("NIM. D021201006").font.name = 'Arial'

    doc.add_page_break()

    # 2.4 PRAKATA (Halaman iv)
    add_heading_1(doc, "PRAKATA")

    add_styled_paragraph(doc, "Puji dan syukur ke hadirat Tuhan Yang Maha Esa atas limpahan rahmat, taufik, dan hidayah-Nya, sehingga penulis dapat menyelesaikan naskah Proposal Tugas Akhir ini dengan judul *\"Analisis Kinematika, Dinamika, dan Estimasi Keadaan Optimal Kalman Filter untuk Vision-Based Tracking pada Over-Actuated 8-Thruster 6-DOF Vectored AUV\"*. Naskah ini diajukan sebagai salah satu syarat akademis kurikuler wajib guna mencapai derajat Sarjana Teknik (S.T.) pada Departemen Teknik Mesin, Fakultas Teknik, Universitas Hasanuddin.")
    add_styled_paragraph(doc, "Penulis menyadari sepenuhnya bahwa kelancaran dan keterwujudan penyusunan proposal ini tidak lepas dari bimbingan, arahan saintifik, motivasi, serta bantuan moril maupun materil dari berbagai pihak. Oleh karena itu, dengan kerendahan hati penulis menyampaikan apresiasi dan ucapan terima kasih yang setinggi-tingginya kepada:")

    ack_list = [
        ("1.", "Bapak Dr. Muhammad Syahid, S.T., M.T., selaku Ketua Departemen Teknik Mesin Fakultas Teknik Universitas Hasanuddin atas segala dukungan akademis, legalitas administrasi, dan penyediaan atmosfer penelitian mekatronika yang unggul."),
        ("2.", "Bapak Andi Amijoyo Mochtar, S.T., M.Sc., Ph.D., selaku Dosen Pembimbing Utama atas curahan bimbingan intelektual, arahan matematis, ketelitian permodelan dinamika wahana, serta dorongan dedikatif yang tak henti-hentinya selama riset ini diformulasikan."),
        ("3.", "Seluruh Dosen dan Staf Pengajar Departemen Teknik Mesin Universitas Hasanuddin yang telah mentransfer khazanah keilmuan keteknikan yang berharga selama masa studi sarjana penulis."),
        ("4.", "Kedua orang tua tercinta dan segenap keluarga besar penulis atas untaian doa yang tiada putus, pengorbanan materil, ketulusan cinta, dan dukungan moril yang menjadi sumber energi terbesar bagi penulis."),
        ("5.", "Rekan-rekan peneliti dan asisten di Laboratorium Mekatronika dan Robotika, Departemen Teknik Mesin, serta segenap rekan mahasiswa Departemen Teknik Mesin Universitas Hasanuddin Angkatan 2020 atas sinergi diskusi laboratorium, bantuan pengujian eksperimental, dan kebersamaan yang senantiasa terbangun.")
    ]

    for num, txt in ack_list:
        p_ack = doc.add_paragraph()
        p_ack.paragraph_format.left_indent = Mm(10)
        p_ack.paragraph_format.first_line_indent = Mm(-6)
        p_ack.paragraph_format.space_after = Pt(4)
        p_ack.paragraph_format.line_spacing = 1.15
        p_ack.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_num = p_ack.add_run(num + "  ")
        r_num.bold = True
        r_num.font.name = 'Arial'
        for t_type, t_val in parse_inline_runs(txt):
            r_tx = p_ack.add_run(t_val)
            r_tx.font.name = 'Arial'
            if t_type == 'italic':
                r_tx.italic = True
            if t_type == 'bold':
                r_tx.bold = True

    add_styled_paragraph(doc, "Penulis menyadari bahwa naskah proposal ini masih memiliki keterbatasan. Kritik dan saran yang membangun sangat penulis harapkan guna penyempurnaan implementasi penelitian di masa mendatang.")

    p_sig2 = doc.add_paragraph()
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig2.paragraph_format.space_before = Pt(16)
    p_sig2.paragraph_format.line_spacing = 1.15
    p_sig2.add_run("Makassar, 24 September 2026\n\n\n").font.name = 'Arial'
    r_pn = p_sig2.add_run("Penulis")
    r_pn.bold = True
    r_pn.font.name = 'Arial'

    doc.add_page_break()

    # 2.5 ABSTRAK (INDONESIAN - Halaman v) - No English abstract per user request
    add_heading_1(doc, "ABSTRAK")

    p_abs_title = doc.add_paragraph()
    p_abs_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs_title.paragraph_format.space_after = Pt(8)
    p_abs_title.paragraph_format.line_spacing = 1.15
    for t_type, t_val in parse_inline_runs(OFFICIAL_TITLE_MD):
        r_at = p_abs_title.add_run(t_val)
        r_at.font.name = 'Arial'
        r_at.font.size = Pt(10)
        r_at.bold = True
        if t_type == 'italic':
            r_at.italic = True

    abstract_id_text = (
        "Latar Belakang. Wahana bawah air nirawak *Autonomous Underwater Vehicle* (AUV) konvensional dengan konfigurasi *underactuated* "
        "memiliki keterbatasan kendali orientasi ruang, terutama kopling hidrodinamika antara gerak *surge*, *heave*, dan *pitch*. "
        "Tujuan. Penelitian ini bertujuan merumuskan pemodelan komprehensif kinematika dan dinamika 6 *degrees of freedom* (6-DOF) "
        "benda tegar bawah air (*rigid-body*) Fossen, merancang matriks alokasi gaya dorong (*thrust allocation*) 8 motor *thruster* *brushless*, "
        "serta menerapkan penapis *Extended Kalman Filter* (EKF) adaptif terdistribusi untuk estimasi keadaan visual dan gerak wahana. "
        "Metode. Pemodelan sistem mengintegrasikan tensor massa inersia dan massa tambah hidrodinamika ($$\\mathbf{M} = \\mathbf{M}_{RB} + \\mathbf{M}_A$$), "
        "matriks redaman linier-kuadratik $$\\mathbf{D}(\\boldsymbol{\\nu}_r)$$, gaya pemulih hidrostatis $$\\mathbf{g}(\\boldsymbol{\\eta})$$, dan alokasi gaya dorong matriks *pseudo-inverse* "
        "Moore-Penrose $$\\mathbf{T}^\\dagger_{8 \\times 6}$$. Estimasi pelacakan target visual memanfaatkan *Discrete Kalman Filter* berorde 8 pada *bounding box* YOLO "
        "dengan *Continuous White Noise Acceleration* (CWNA) dan *outlier gating* Mahalanobis ($$D_M^2 \\le 9{,}488$$), serta EKF navigasi inersial terdistribusi pada *companion computer* Raspberry Pi 4B. "
        "Arsitektur divalidasi melalui skema *Software-In-The-Loop* (SITL) Gazebo Harmonic + ArduSub dan *Hardware-In-The-Loop* (HITL) Pixhawk. "
        "Hasil yang Diharapkan. Penelitian menghasilkan rumusan analitis matriks alokasi dorongan yang stabil pada mode *pitch-holding* aktif, "
        "mereduksi kesalahan estimasi titik tengah target visual sebesar >40% terhadap derau oklusi, serta mempertahankan kestabilan sikap 6-DOF secara *real-time*. "
        "Kesimpulan. Pendekatan terpadu antara hidrodinamika 6-DOF dan penapis optimal Kalman menyediakan landasan saintifik dan mekatronika yang tangguh untuk misi inspeksi bawah air otonom."
    )

    p_abs_body = doc.add_paragraph()
    p_abs_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs_body.paragraph_format.space_after = Pt(8)
    p_abs_body.paragraph_format.line_spacing = 1.0
    for t_type, t_val in parse_inline_runs(abstract_id_text):
        r_ab = p_abs_body.add_run(t_val)
        r_ab.font.name = 'Arial'
        r_ab.font.size = Pt(9.5)
        if t_type == 'italic':
            r_ab.italic = True
        if t_type == 'bold':
            r_ab.bold = True

    p_kwi = doc.add_paragraph()
    p_kwi.paragraph_format.space_after = Pt(14)
    p_kwi.paragraph_format.line_spacing = 1.0
    r_kwi_lbl = p_kwi.add_run("Kata Kunci: ")
    r_kwi_lbl.bold = True
    r_kwi_lbl.font.name = 'Arial'
    r_kwi_lbl.font.size = Pt(9.5)
    for t_type, t_val in parse_inline_runs("AUV *over-actuated*; dinamika 6-DOF; alokasi gaya dorong; *Extended Kalman Filter*; *vision-based tracking*; *Hardware-In-The-Loop*."):
        r_kv = p_kwi.add_run(t_val)
        r_kv.font.name = 'Arial'
        r_kv.font.size = Pt(9.5)
        if t_type == 'italic':
            r_kv.italic = True

    doc.add_page_break()

    # 2.6 DAFTAR ISI (Halaman vi)
    add_heading_1(doc, "DAFTAR ISI")

    p_th = doc.add_paragraph()
    p_th.paragraph_format.tab_stops.add_tab_stop(PRINTABLE_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p_th.paragraph_format.space_after = Pt(8)
    r_th1 = p_th.add_run("Judul")
    r_th1.bold = True
    r_th1.font.name = 'Arial'
    r_th1.font.size = Pt(9.5)
    r_th2 = p_th.add_run("\tHalaman")
    r_th2.bold = True
    r_th2.font.name = 'Arial'
    r_th2.font.size = Pt(9.5)

    add_leader_line(doc, "HALAMAN SAMPUL DEPAN", "i", bold=True)
    add_leader_line(doc, "LEMBAR PENGESAHAN PROPOSAL TUGAS AKHIR", "ii", bold=True)
    add_leader_line(doc, "PERNYATAAN KEASLIAN PROPOSAL TUGAS AKHIR", "iii", bold=True)
    add_leader_line(doc, "PRAKATA", "iv", bold=True)
    add_leader_line(doc, "ABSTRAK", "v", bold=True)
    add_leader_line(doc, "DAFTAR ISI", "vi", bold=True)
    add_leader_line(doc, "DAFTAR TABEL", "vii", bold=True)
    add_leader_line(doc, "DAFTAR GAMBAR", "viii", bold=True)
    add_leader_line(doc, "DAFTAR SINGKATAN, ISTILAH, DAN LAMBANG", "ix", bold=True)

    add_leader_line(doc, "BAB I: PENDAHULUAN", "1", bold=True, space_after=4)
    add_leader_line(doc, "1.1 Latar Belakang", "1", indent_mm=5)
    add_leader_line(doc, "1.2 Rumusan Masalah", "5", indent_mm=5)
    add_leader_line(doc, "1.3 Tujuan Penelitian", "6", indent_mm=5)
    add_leader_line(doc, "1.4 Batasan Masalah", "7", indent_mm=5)
    add_leader_line(doc, "1.5 Manfaat Penelitian", "8", indent_mm=5)

    add_leader_line(doc, "BAB II: TINJAUAN PUSTAKA", "10", bold=True, space_after=4)
    add_leader_line(doc, "2.1 Tinjauan Pustaka (*State of the Art* Penelitian AUV)", "10", indent_mm=5)
    add_leader_line(doc, "2.2 Sistem Koordinat dan Konvensi SNAME", "13", indent_mm=5)
    add_leader_line(doc, "2.3 Penurunan Kinematika 6-DOF dan Matriks Jacobian", "18", indent_mm=5)
    add_leader_line(doc, "2.4 Penurunan Dinamika Hidrodinamika 6-DOF (Persamaan Fossen)", "24", indent_mm=5)
    add_leader_line(doc, "2.5 Alokasi Gaya Dorong Sistem *Over-Actuated* 8-Pendorong", "32", indent_mm=5)
    add_leader_line(doc, "2.6 Teori dan Formulasi Optimal *Kalman Filter* Suite", "37", indent_mm=5)

    add_leader_line(doc, "BAB III: METODOLOGI PENELITIAN", "52", bold=True, space_after=4)
    add_leader_line(doc, "3.1 Tempat dan Waktu Penelitian", "52", indent_mm=5)
    add_leader_line(doc, "3.2 Diagram Alir Penelitian", "54", indent_mm=5)
    add_leader_line(doc, "3.3 Identifikasi Parameter Fisik dan Hidrodinamika Wahana", "56", indent_mm=5)
    add_leader_line(doc, "3.4 Perancangan Arsitektur *Software-In-The-Loop* (SITL)", "61", indent_mm=5)
    add_leader_line(doc, "3.5 Perancangan Arsitektur *Hardware-In-The-Loop* (HITL)", "64", indent_mm=5)
    add_leader_line(doc, "3.6 Prosedur Pengujian dan Evaluasi Kinerja", "71", indent_mm=5)

    add_leader_line(doc, "DAFTAR PUSTAKA", "76", bold=True, space_after=4)

    doc.add_page_break()

    # 2.7 DAFTAR TABEL (Halaman vii) - Exactly matches the 8 actual tables present in the thesis
    add_heading_1(doc, "DAFTAR TABEL")

    p_tth = doc.add_paragraph()
    p_tth.paragraph_format.tab_stops.add_tab_stop(PRINTABLE_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p_tth.paragraph_format.space_after = Pt(8)
    r_tth1 = p_tth.add_run("Nomor Urut dan Judul Tabel")
    r_tth1.bold = True
    r_tth1.font.name = 'Arial'
    r_tth1.font.size = Pt(9.5)
    r_tth2 = p_tth.add_run("\tHalaman")
    r_tth2.bold = True
    r_tth2.font.name = 'Arial'
    r_tth2.font.size = Pt(9.5)

    tables_info = [
        ("Tabel 2.1", "Matriks Sintesis Literatur Terkini (2021–2025) Bidang Dinamika dan Kontrol AUV", "11"),
        ("Tabel 2.2", "Notasi dan Konvensi 6 Derajat Kebebasan SNAME (1950) & Fossen (2021)", "15"),
        ("Tabel 2.3", "Koordinat Spasial dan Vektor Orientasi 8 Pendorong Wahana Over-Actuated", "34"),
        ("Tabel 3.1", "Parameter Fisik dan Properti Benda Tegar Wahana Over-Actuated 8-Pendorong", "57"),
        ("Tabel 3.2", "Koefisien Derivatif Massa Tambah Hidrodinamika Wahana", "58"),
        ("Tabel 3.3", "Koefisien Redaman Hidrodinamika Linier dan Kuadratik Wahana", "59"),
        ("Tabel 3.4", "Posisi Spasial dan Vektor Satuan Gaya Dorong 8-Pendorong Bervektor", "60"),
        ("Tabel 3.5", "Spesifikasi Komponen Perangkat Keras Arsitektur HITL", "64")
    ]

    for num, title, pg in tables_info:
        add_leader_line(doc, f"{num}  {title}", pg, space_after=3)

    doc.add_page_break()

    # 2.8 DAFTAR GAMBAR (Halaman viii) - Exactly matches the 10 actual figures present in the thesis
    add_heading_1(doc, "DAFTAR GAMBAR")

    p_fgh = doc.add_paragraph()
    p_fgh.paragraph_format.tab_stops.add_tab_stop(PRINTABLE_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p_fgh.paragraph_format.space_after = Pt(8)
    r_fgh1 = p_fgh.add_run("Nomor Urut dan Judul Gambar")
    r_fgh1.bold = True
    r_fgh1.font.name = 'Arial'
    r_fgh1.font.size = Pt(9.5)
    r_fgh2 = p_fgh.add_run("\tHalaman")
    r_fgh2.bold = True
    r_fgh2.font.name = 'Arial'
    r_fgh2.font.size = Pt(9.5)

    figures_info = [
        ("Gambar 2.1", "Sistem Kerangka Acuan Inersia Bumi (Fn - NED) dan Kerangka Acuan Bergerak Bodi (Fb - FRD) Konvensi SNAME (1950) dan Fossen (2021)", "13"),
        ("Gambar 3.1", "Diagram Alir Tahapan Penelitian Komprehensif", "54"),
        ("Gambar 3.2", "Arsitektur Simulasi Software-In-The-Loop (SITL) Sistem AUV", "62"),
        ("Gambar 3.3", "Arsitektur Integrasi Hardware-In-The-Loop (HITL) Mekatronika AUV", "64"),
        ("Gambar 3.4", "Rangka (Frame) dan Lambung Tekanan Kustom AUV 8-Pendorong", "65"),
        ("Gambar 3.5", "Papan Pengendali Penerbangan (Flight Controller) Pixhawk 2.4.8", "65"),
        ("Gambar 3.6", "Komputer Pendamping (Companion Computer) Raspberry Pi 4B", "66"),
        ("Gambar 3.7", "Modul Pengendali Kecepatan Elektronik (ESC EMAX BLHeli 30A)", "66"),
        ("Gambar 3.8", "Motor Pendorong Bawah Air (BLDC Underwater Thruster)", "67"),
        ("Gambar 3.9", "Sumber Daya Baterai Li-Po 4S 14.8V 6000 mAh dan Pengisi Daya SKYRC IMAX B6AC V2", "67")
    ]

    for num, title, pg in figures_info:
        add_leader_line(doc, f"{num}  {title}", pg, space_after=3)

    doc.add_page_break()

    # 2.9 DAFTAR SINGKATAN DAN ISTILAH (Halaman ix)
    add_heading_1(doc, "DAFTAR SINGKATAN, ISTILAH, DAN LAMBANG")
    add_heading_2(doc, "1. Daftar Singkatan dan Akronim")

    abbrevs = [
        ("AUV", "Autonomous Underwater Vehicle (Wahana Bawah Air Otonom)"),
        ("ROV", "Remotely Operated Vehicle (Wahana Bawah Air Kendali Jarak Jauh)"),
        ("DOF", "Degrees of Freedom (Derajat Kebebasan Gerak Spasial)"),
        ("EKF", "Extended Kalman Filter (Penapis Kalman Diperluas)"),
        ("UKF", "Unscented Kalman Filter (Penapis Kalman Tanpa Aroma)"),
        ("CWNA", "Continuous White Noise Acceleration (Akselerasi Derau Putih Kontinu)"),
        ("YOLO", "You Only Look Once (Jaringan Syaraf Deteksi Objek Real-Time)"),
        ("SITL", "Software-In-The-Loop (Simulasi Perangkat Lunak Tertutup)"),
        ("HITL", "Hardware-In-The-Loop (Simulasi Berbantuan Perangkat Keras Riil)"),
        ("NED", "North-East-Down (Sistem Sumbu Koordinat Bumi: Utara-Timur-Bawah)"),
        ("FRD", "Forward-Right-Down (Sistem Sumbu Koordinat Bodi: Depan-Kanan-Bawah)"),
        ("SNAME", "The Society of Naval Architects and Marine Engineers"),
        ("ROS", "Robot Operating System (Middleware Komunikasi Robotika)"),
        ("CLAHE", "Contrast Limited Adaptive Histogram Equalization"),
        ("MAVLink", "Micro Air Vehicle Link (Protokol Serial Telemetri Robotika)"),
        ("PWM", "Pulse Width Modulation (Modulasi Lebar Pulsa Kontrol ESC)"),
        ("ESC", "Electronic Speed Controller (Pengendali Kecepatan Motor Brushless)"),
        ("IMU", "Inertial Measurement Unit (Sensor Akselerometer dan Giroskop Inersial)")
    ]

    t_abb = doc.add_table(rows=len(abbrevs)+1, cols=2)
    t_abb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c in t_abb.rows[0].cells:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EAEFF5"/>')
        c._tc.get_or_add_tcPr().append(shd)
        set_cell_borders(c)
        set_cell_margins(c, top=80, bottom=80, left=120, right=120)

    t_abb.rows[0].cells[0].paragraphs[0].add_run("Singkatan").bold = True
    t_abb.rows[0].cells[1].paragraphs[0].add_run("Kepanjangan / Arti Teknis").bold = True
    t_abb.rows[0].cells[0].width = Mm(28)
    t_abb.rows[0].cells[1].width = Mm(103)

    for r_idx, (abbr, desc) in enumerate(abbrevs):
        row_cells = t_abb.rows[r_idx+1].cells
        set_cell_borders(row_cells[0])
        set_cell_borders(row_cells[1])
        set_cell_margins(row_cells[0], top=60, bottom=60, left=120, right=120)
        set_cell_margins(row_cells[1], top=60, bottom=60, left=120, right=120)
        row_cells[0].width = Mm(28)
        row_cells[1].width = Mm(103)
        p0 = row_cells[0].paragraphs[0]
        p0.paragraph_format.line_spacing = 1.0
        p0.paragraph_format.space_after = Pt(2)
        r_ab = p0.add_run(abbr)
        r_ab.font.name = 'Arial'
        r_ab.font.size = Pt(9)
        r_ab.bold = True

        p1 = row_cells[1].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.0
        p1.paragraph_format.space_after = Pt(2)
        r_de = p1.add_run(desc)
        r_de.font.name = 'Arial'
        r_de.font.size = Pt(9)

    for r_idx, row in enumerate(t_abb.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    p_gap = doc.add_paragraph()
    p_gap.paragraph_format.space_before = Pt(12)

    # 2.10 DAFTAR LAMBANG DAN SIMBOL MATEMATIKA
    add_heading_2(doc, "2. Daftar Lambang dan Simbol Matematika")

    symbols_latex = [
        ("$$\\mathcal{F}^n$$", "-", "Kerangka Acuan Inersia Bumi (*Earth-Fixed NED Frame*) $$\\{O_n, x_n, y_n, z_n\\}$$"),
        ("$$\\mathcal{F}^b$$", "-", "Kerangka Acuan Bergerak Bodi (*Body-Fixed Frame*) $$\\{O_b, x_b, y_b, z_b\\}$$"),
        ("$$\\boldsymbol{\\eta}$$", "$$\\mathbb{R}^6\\text{ (m, rad)}$$", "Vektor posisi dan orientasi spasial di $$\\mathcal{F}^n$$: $$[x, y, z, \\phi, \\theta, \\psi]^T$$"),
        ("$$\\boldsymbol{\\nu}$$", "$$\\mathbb{R}^6\\text{ (m/s, rad/s)}$$", "Vektor kecepatan linier dan sudut di $$\\mathcal{F}^b$$: $$[u, v, w, p, q, r]^T$$"),
        ("$$\\boldsymbol{\\tau}$$", "$$\\mathbb{R}^6\\text{ (N, N}\\cdot\\text{m)}$$", "Vektor gaya dan momen generalisasi di $$\\mathcal{F}^b$$: $$[X, Y, Z, K, M, N]^T$$"),
        ("$$\\boldsymbol{\\nu}_c$$", "$$\\mathbb{R}^6\\text{ (m/s)}$$", "Vektor kecepatan arus laut fluida pada kerangka bodi: $$[u_c, v_c, w_c, 0, 0, 0]^T$$"),
        ("$$\\boldsymbol{\\nu}_r$$", "$$\\mathbb{R}^6\\text{ (m/s, rad/s)}$$", "Vektor kecepatan relatif wahana terhadap fluida: $$\\boldsymbol{\\nu}_r = \\boldsymbol{\\nu} - \\boldsymbol{\\nu}_c$$"),
        ("$$\\mathbf{R}_b^n(\\boldsymbol{\\eta}_2)$$", "$$SO(3)$$", "Matriks transformasi rotasi ortogonal dari $$\\mathcal{F}^b$$ ke $$\\mathcal{F}^n$$"),
        ("$$\\mathbf{T}_\\Theta(\\boldsymbol{\\eta}_2)$$", "$$\\mathbb{R}^{3 \\times 3}$$", "Matriks transformasi kecepatan sudut Euler: $$\\dot{\\boldsymbol{\\eta}}_2 = \\mathbf{T}_\\Theta \\boldsymbol{\\nu}_2$$"),
        ("$$\\mathbf{J}(\\boldsymbol{\\eta}_2)$$", "$$\\mathbb{R}^{6 \\times 6}$$", "Matriks Jacobian kinematika gabungan: $$\\text{diag}[\\mathbf{R}_b^n, \\mathbf{T}_\\Theta]$$"),
        ("$$\\mathbf{q}$$", "$$S^3$$", "Kuaternion unit orientasi empat-dimensi: $$[\\eta, \\epsilon_1, \\epsilon_2, \\epsilon_3]^T$$"),
        ("$$\\mathbf{M}_{RB}$$", "$$\\mathbb{R}^{6 \\times 6}\\text{ (kg, kg}\\cdot\\text{m}^2)$$", "Tensor massa inersia bodi kaku (*rigid-body mass matrix*)"),
        ("$$\\mathbf{M}_A$$", "$$\\mathbb{R}^{6 \\times 6}\\text{ (kg, kg}\\cdot\\text{m}^2)$$", "Tensor massa tambah hidrodinamika fluida (*hydrodynamic added mass*)"),
        ("$$\\mathbf{M}$$", "$$\\mathbb{R}^{6 \\times 6}\\text{ (kg, kg}\\cdot\\text{m}^2)$$", "Tensor massa sistem total gabungan: $$\\mathbf{M} = \\mathbf{M}_{RB} + \\mathbf{M}_A$$"),
        ("$$\\mathbf{C}_{RB}(\\boldsymbol{\\nu})$$", "$$\\mathbb{R}^{6 \\times 6}\\text{ (N}\\cdot\\text{s/m, N}\\cdot\\text{s}\\cdot\\text{m)}$$", "Matriks gaya Coriolis dan sentripetal bodi kaku"),
        ("$$\\mathbf{C}_A(\\boldsymbol{\\nu}_r)$$", "$$\\mathbb{R}^{6 \\times 6}\\text{ (N}\\cdot\\text{s/m, N}\\cdot\\text{s}\\cdot\\text{m)}$$", "Matriks Coriolis dan sentripetal massa tambah hidrodinamika"),
        ("$$\\mathbf{D}(\\boldsymbol{\\nu}_r)$$", "$$\\mathbb{R}^{6 \\times 6}\\text{ (N}\\cdot\\text{s/m, N}\\cdot\\text{s}^2/\\text{m}^2)$$", "Tensor redaman hidrodinamika gabungan (linier $$\\mathbf{D}_L$$ + kuadratik $$\\mathbf{D}_{NL}$$)"),
        ("$$\\mathbf{g}(\\boldsymbol{\\eta})$$", "$$\\mathbb{R}^6\\text{ (N, N}\\cdot\\text{m)}$$", "Vektor gaya dan momen pemulih hidrostatis (gravitasi dan gaya apung)"),
        ("$$GM_T$$", "$$\\text{m}$$", "Tinggi metasentris transversal wahana: $$z_g - z_b$$"),
        ("$$\\mathbf{T}_{6 \\times 8}$$", "$$\\mathbb{R}^{6 \\times 8}$$", "Matriks konfigurasi geometri dan alokasi gaya dorong 8 motor pendorong"),
        ("$$\\mathbf{T}^\\dagger$$", "$$\\mathbb{R}^{8 \\times 6}$$", "Matriks *pseudo-inverse* Moore-Penrose: $$\\mathbf{T}^T(\\mathbf{T}\\mathbf{T}^T)^{-1}$$"),
        ("$$\\mathbf{f}$$", "$$\\mathbb{R}^{8}\\text{ (N)}$$", "Vektor gaya dorong individual 8 motor pendorong: $$[f_1, f_2, \\dots, f_8]^T$$"),
        ("$$\\mathbf{x}_k$$", "$$\\mathbb{R}^8\\text{ (px, px/s)}$$", "Vektor keadaan penjejakan visual: $$[x, y, w, h, v_x, v_y, v_w, v_h]^T$$"),
        ("$$\\mathbf{P}_k$$", "$$\\mathbb{R}^{n \\times n}$$", "Matriks kovariansi kesalahan estimasi filter (*error covariance matrix*)"),
        ("$$\\mathbf{K}_k$$", "-", "Matriks penguatan optimal Kalman (*optimal Kalman gain*)"),
        ("$$\\mathbf{Q}$$", "$$\\text{px}^2/\\text{s}^3$$", "Matriks kovariansi derau proses (*process noise covariance matrix*)"),
        ("$$\\mathbf{R}$$", "$$\\text{px}^2$$", "Matriks kovariansi derau pengukuran (*measurement noise covariance matrix*)"),
        ("$$D_M^2$$", "-", "Jarak kuadratis Mahalanobis untuk validasi inovasi dan *outlier gating*")
    ]

    t_sym = doc.add_table(rows=len(symbols_latex)+1, cols=3)
    t_sym.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c in t_sym.rows[0].cells:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EAEFF5"/>')
        c._tc.get_or_add_tcPr().append(shd)
        set_cell_borders(c)
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)

    t_sym.rows[0].cells[0].paragraphs[0].add_run("Simbol").bold = True
    t_sym.rows[0].cells[1].paragraphs[0].add_run("Dimensi / Satuan").bold = True
    t_sym.rows[0].cells[2].paragraphs[0].add_run("Definisi Matematis dan Fisik").bold = True
    t_sym.rows[0].cells[0].width = Mm(26)
    t_sym.rows[0].cells[1].width = Mm(32)
    t_sym.rows[0].cells[2].width = Mm(73)

    for r_idx, (sym, dim, defn) in enumerate(symbols_latex):
        row_cells = t_sym.rows[r_idx+1].cells
        set_cell_borders(row_cells[0])
        set_cell_borders(row_cells[1])
        set_cell_borders(row_cells[2])
        set_cell_margins(row_cells[0], top=50, bottom=50, left=100, right=100)
        set_cell_margins(row_cells[1], top=50, bottom=50, left=100, right=100)
        set_cell_margins(row_cells[2], top=50, bottom=50, left=100, right=100)

        row_cells[0].width = Mm(26)
        row_cells[1].width = Mm(32)
        row_cells[2].width = Mm(73)

        p0 = row_cells[0].paragraphs[0]
        p0.paragraph_format.line_spacing = 1.0
        p0.paragraph_format.space_after = Pt(2)
        for t_type, t_val in parse_inline_runs(sym):
            r = p0.add_run(t_val)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.bold = True

        p1 = row_cells[1].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.0
        p1.paragraph_format.space_after = Pt(2)
        for t_type, t_val in parse_inline_runs(dim):
            r = p1.add_run(t_val)
            r.font.name = 'Arial'
            r.font.size = Pt(8.5)

        p2 = row_cells[2].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.0
        p2.paragraph_format.space_after = Pt(2)
        for t_type, t_val in parse_inline_runs(defn):
            r = p2.add_run(t_val)
            r.font.name = 'Arial'
            r.font.size = Pt(8.5)
            if t_type == 'italic':
                r.italic = True
            if t_type == 'bold':
                r.bold = True

    for r_idx, row in enumerate(t_sym.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def process_markdown_chapter(doc, md_filepath, chapter_title_override=None):
    with open(md_filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    table_rows = []
    in_code_block = False
    code_lines = []
    first_h1_handled = False

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                if code_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Mm(10)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.0
                    r = p.add_run('\n'.join(code_lines))
                    r.font.name = 'Consolas'
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Markdown tables
        if stripped.startswith('|') and stripped.endswith('|'):
            if re.match(r'^\|[\s\-:]+(\|[\s\-:]+)+\|$', stripped):
                i += 1
                continue
            cells = [c.strip() for c in stripped[1:-1].split('|')]
            table_rows.append(cells)
            i += 1
            if i >= len(lines) or not (lines[i].strip().startswith('|') and lines[i].strip().endswith('|')):
                if table_rows:
                    n_cols = max(len(r) for r in table_rows)
                    t = doc.add_table(rows=len(table_rows), cols=n_cols)
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER

                    col_w = PRINTABLE_WIDTH / n_cols
                    for r_idx, r_data in enumerate(table_rows):
                        for c_idx in range(n_cols):
                            cell = t.rows[r_idx].cells[c_idx]
                            cell.width = col_w
                            set_cell_borders(cell)
                            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                            val = r_data[c_idx] if c_idx < len(r_data) else ""
                            if r_idx == 0:
                                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EAEFF5"/>')
                                cell._tc.get_or_add_tcPr().append(shd)
                            p = cell.paragraphs[0]
                            p.paragraph_format.line_spacing = 1.0
                            p.paragraph_format.space_after = Pt(2)
                            tokens = parse_inline_runs(val)
                            for t_type, t_val in tokens:
                                r = p.add_run(t_val)
                                r.font.name = 'Arial'
                                r.font.size = Pt(8.5)
                                if r_idx == 0:
                                    r.bold = True
                                else:
                                    r.bold = False
                                if t_type == 'italic':
                                    r.italic = True
                    for r_idx, row in enumerate(t.rows):
                        trPr = row._tr.get_or_add_trPr()
                        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                        if r_idx == 0:
                            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                    table_rows = []
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if not stripped:
            i += 1
            continue
        if stripped in ['---', '***', '___']:
            i += 1
            continue
        if stripped.startswith('<!--') and stripped.endswith('-->'):
            i += 1
            continue
        if stripped.startswith('<div') or stripped.startswith('</div'):
            i += 1
            continue
        if stripped.startswith('## ANALISIS KINEMATIKA, DINAMIKA') or stripped.startswith('# ANALISIS KINEMATIKA'):
            i += 1
            continue
        if stripped.startswith('> **Catatan Format') or stripped.startswith('> **Author') or stripped.startswith('> **Global Synchronized'):
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            if not first_h1_handled and chapter_title_override:
                r = p.add_run(chapter_title_override)
                r.bold = True
                r.font.name = 'Arial'
                r.font.size = Pt(11)
            else:
                for t_type, t_val in parse_inline_runs(h_text):
                    r = p.add_run(t_val)
                    r.font.name = 'Arial'
                    r.font.size = Pt(11)
                    r.bold = True
                    if t_type == 'italic':
                        r.italic = True
            first_h1_handled = True
            i += 1
            continue

        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            for t_type, t_val in parse_inline_runs(h_text):
                r = p.add_run(t_val)
                r.font.name = 'Arial'
                r.font.size = Pt(10)
                r.bold = True
                if t_type == 'italic':
                    r.italic = True
            i += 1
            continue

        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            for t_type, t_val in parse_inline_runs(h_text):
                r = p.add_run(t_val)
                r.font.name = 'Arial'
                r.font.size = Pt(10)
                r.bold = True
                if t_type == 'italic':
                    r.italic = True
            i += 1
            continue

        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            for t_type, t_val in parse_inline_runs(h_text):
                r = p.add_run(t_val)
                r.font.name = 'Arial'
                r.font.size = Pt(10)
                r.bold = True
                if t_type == 'italic':
                    r.italic = True
            i += 1
            continue

        # Bullet lists and citations
        if (stripped.startswith('- ') or stripped.startswith('* ')) and not stripped.startswith('***'):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15

            # Check if this item is a citation entry (e.g., *[1]*, [1], etc.)
            m_cite = re.match(r'^\*?\[\d+\]\*?\s*', item_text)
            if m_cite:
                p.paragraph_format.left_indent = Mm(10)
                p.paragraph_format.first_line_indent = Mm(-8)
            else:
                p.paragraph_format.left_indent = Mm(10)
                p.paragraph_format.first_line_indent = Mm(-5)
                r_b = p.add_run("•  ")
                r_b.font.name = 'Arial'
                r_b.font.size = Pt(10)
                r_b.bold = False

            for t_type, t_val in parse_inline_runs(item_text):
                r = p.add_run(t_val)
                r.font.name = 'Arial'
                r.font.size = Pt(10)
                r.bold = False
                if t_type == 'italic':
                    r.italic = True
                if t_type == 'code':
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
            i += 1
            continue

        # Numbered lists
        m_num = re.match(r'^(\d+\.)\s+(.*)', stripped)
        if m_num:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(10)
            p.paragraph_format.first_line_indent = Mm(-6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r_num = p.add_run(m_num.group(1) + "  ")
            r_num.font.name = 'Arial'
            r_num.font.size = Pt(10)
            r_num.bold = False
            for t_type, t_val in parse_inline_runs(m_num.group(2)):
                r = p.add_run(t_val)
                r.font.name = 'Arial'
                r.font.size = Pt(10)
                r.bold = False
                if t_type == 'italic':
                    r.italic = True
                if t_type == 'code':
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
            i += 1
            continue

        # Blockquotes
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            for t_type, t_val in parse_inline_runs(stripped[2:].strip()):
                r = p.add_run(t_val)
                r.font.name = 'Arial'
                r.font.size = Pt(9.5)
                r.italic = True
                r.bold = False
            i += 1
            continue

        # Images - Centered alignment
        m_img = re.match(r'^!\[(.*?)\]\((.*?)\)', stripped)
        if m_img:
            alt_text = m_img.group(1).strip()
            raw_path = m_img.group(2).strip()

            resolved_path = None
            candidate_paths = [
                raw_path,
                os.path.join(thesis_dir, raw_path),
                os.path.join(thesis_dir, "figures", os.path.basename(raw_path)),
                os.path.join("/home/radhi/Documents/AUV Development/Thesis", raw_path),
                os.path.join("/home/radhi/Documents/AUV Development/Thesis/figures", os.path.basename(raw_path))
            ]
            for cp in candidate_paths:
                if os.path.isfile(cp):
                    resolved_path = cp
                    break

            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)

            if resolved_path:
                r_img = p_img.add_run()
                target_width = Mm(128)
                try:
                    from PIL import Image as PILImage
                    with PILImage.open(resolved_path) as im:
                        w_px, h_px = im.size
                        aspect = h_px / w_px
                        if aspect > 0.8:
                            target_width = Mm(105)
                except Exception:
                    pass
                r_img.add_picture(resolved_path, width=target_width)
            else:
                p_img.add_run("[").font.name = 'Arial'
                for t_type, t_val in parse_inline_runs(alt_text):
                    r_c = p_img.add_run(t_val)
                    r_c.font.name = 'Arial'
                    r_c.font.size = Pt(9.5)
                    r_c.italic = True
                    r_c.bold = False
                p_img.add_run("]").font.name = 'Arial'
            i += 1
            continue

        # Table caption - strictly at TOP, centered, no trailing period
        if stripped.startswith('**Tabel ') or stripped.startswith('Tabel ') or stripped.startswith('*Tabel '):
            p_tt = doc.add_paragraph()
            p_tt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_tt.paragraph_format.space_before = Pt(8)
            p_tt.paragraph_format.space_after = Pt(3)
            p_tt.paragraph_format.line_spacing = 1.15
            m_tcap = re.match(r'^(\*?\*?Tabel\s+[\d\.]+\*?\*?)\s*(.*)', stripped)
            if m_tcap:
                prefix = m_tcap.group(1).replace('*', '').replace('#', '').strip()
                title_part = m_tcap.group(2).strip().rstrip('.')
                r1 = p_tt.add_run(prefix + " ")
                r1.font.name = 'Arial'
                r1.font.size = Pt(9.5)
                r1.bold = True
                for t_type, t_val in parse_inline_runs(title_part):
                    r2 = p_tt.add_run(t_val)
                    r2.font.name = 'Arial'
                    r2.font.size = Pt(9.5)
                    r2.bold = False
                    if t_type == 'italic':
                        r2.italic = True
            else:
                for t_type, t_val in parse_inline_runs(stripped.rstrip('.')):
                    r_tt = p_tt.add_run(t_val)
                    r_tt.font.name = 'Arial'
                    r_tt.font.size = Pt(9.5)
                    r_tt.bold = False
                    if t_type == 'italic':
                        r_tt.italic = True
            i += 1
            continue

        # Figure caption - strictly at BOTTOM, centered, no trailing period
        if stripped.startswith('**Gambar ') or stripped.startswith('Gambar ') or stripped.startswith('*Gambar '):
            p_fc = doc.add_paragraph()
            p_fc.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_fc.paragraph_format.space_before = Pt(4)
            p_fc.paragraph_format.space_after = Pt(8)
            p_fc.paragraph_format.line_spacing = 1.15
            m_fcap = re.match(r'^(\*?\*?Gambar\s+[\d\.]+\*?\*?)\s*(.*)', stripped)
            if m_fcap:
                prefix = m_fcap.group(1).replace('*', '').replace('#', '').strip()
                title_part = m_fcap.group(2).strip().rstrip('.')
                r1 = p_fc.add_run(prefix + " ")
                r1.font.name = 'Arial'
                r1.font.size = Pt(9.5)
                r1.bold = True
                for t_type, t_val in parse_inline_runs(title_part):
                    r2 = p_fc.add_run(t_val)
                    r2.font.name = 'Arial'
                    r2.font.size = Pt(9.5)
                    r2.bold = False
                    if t_type == 'italic':
                        r2.italic = True
            else:
                for t_type, t_val in parse_inline_runs(stripped.rstrip('.')):
                    r_fc = p_fc.add_run(t_val)
                    r_fc.font.name = 'Arial'
                    r_fc.font.size = Pt(9.5)
                    r_fc.bold = False
                    if t_type == 'italic':
                        r_fc.italic = True
            i += 1
            continue

        # Centered display equations
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            p_m = doc.add_paragraph()
            p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_m.paragraph_format.space_before = Pt(6)
            p_m.paragraph_format.space_after = Pt(6)
            r_m = p_m.add_run(stripped)
            r_m.font.name = 'Arial'
            r_m.font.size = Pt(10)
            i += 1
            continue

        # Regular text paragraph
        add_styled_paragraph(doc, stripped)
        i += 1

print("--- 1. BUILDING MASTER PROPOSAL DOCX ---")
master_doc = docx.Document()
setup_unhas_section(master_doc, is_front_matter=True)
build_front_matter(master_doc)

# Add section for main body (Bab I onwards)
sec_main = master_doc.add_section(docx.enum.section.WD_SECTION_START.NEW_PAGE)
sec_main.page_width = Mm(176)
sec_main.page_height = Mm(250)
sec_main.top_margin = Mm(22.5)
sec_main.bottom_margin = Mm(22.5)
sec_main.left_margin = Mm(22.5)
sec_main.right_margin = Mm(22.5)
sec_main.header.is_linked_to_previous = False
sec_main.footer.is_linked_to_previous = False
sec_main.different_first_page_header_footer = False

sectPr_main = sec_main._sectPr
sectPr_main.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal" w:start="1"/>'))

p_hdr_main = sec_main.header.paragraphs[0]
p_hdr_main.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_hdr_main.paragraph_format.space_after = Pt(0)
p_hdr_main.paragraph_format.line_spacing = 1.0
r_hm = p_hdr_main.add_run()
r_hm.font.name = 'Arial'
r_hm.font.size = Pt(9.5)
fld_xml_main = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
p_hdr_main._p.append(fld_xml_main)

print("Appending BAB I...")
process_markdown_chapter(master_doc, os.path.join(thesis_dir, "BAB_1_PENDAHULUAN.md"), "BAB I\nPENDAHULUAN")
master_doc.add_page_break()

print("Appending BAB II...")
process_markdown_chapter(master_doc, os.path.join(thesis_dir, "BAB_2_LANDASAN_TEORI.md"), "BAB II\nTINJAUAN PUSTAKA")
master_doc.add_page_break()

print("Appending BAB III...")
process_markdown_chapter(master_doc, os.path.join(thesis_dir, "BAB_3_METODOLOGI_PENELITIAN.md"), "BAB III\nMETODOLOGI PENELITIAN")
master_doc.add_page_break()

print("Appending MASTER BIBLIOGRAPHY...")
process_markdown_chapter(master_doc, os.path.join(thesis_dir, "MASTER_BIBLIOGRAPHY.md"), "DAFTAR PUSTAKA")

master_out = os.path.join(docx_dir, "PROPOSAL_LENGKAP_AUV_RADHI_SHAFEEQ.docx")
master_doc.save(master_out)
print(f"Master docx saved: {master_out} ({os.path.getsize(master_out):,} bytes)")

print("\n--- 2. BUILDING INDIVIDUAL CHAPTER DOCX FILES ---")
chapters = [
    ("BAGIAN_AWAL_PROPOSAL.docx", None, build_front_matter, True, 1),
    ("BAB_1_PENDAHULUAN.docx", os.path.join(thesis_dir, "BAB_1_PENDAHULUAN.md"), "BAB I\nPENDAHULUAN", False, 1),
    ("BAB_2_LANDASAN_TEORI.docx", os.path.join(thesis_dir, "BAB_2_LANDASAN_TEORI.md"), "BAB II\nTINJAUAN PUSTAKA", False, 10),
    ("BAB_3_METODOLOGI_PENELITIAN.docx", os.path.join(thesis_dir, "BAB_3_METODOLOGI_PENELITIAN.md"), "BAB III\nMETODOLOGI PENELITIAN", False, 52),
    ("MASTER_BIBLIOGRAPHY.docx", os.path.join(thesis_dir, "MASTER_BIBLIOGRAPHY.md"), "DAFTAR PUSTAKA", False, 76)
]

for out_name, md_file, extra, is_fm, start_pg in chapters:
    c_doc = docx.Document()
    setup_unhas_section(c_doc, is_front_matter=is_fm, start_page=start_pg)
    if extra and callable(extra):
        extra(c_doc)
    elif md_file:
        process_markdown_chapter(c_doc, md_file, extra)
    c_out = os.path.join(docx_dir, out_name)
    c_doc.save(c_out)
    print(f"Chapter docx saved: {c_out} ({os.path.getsize(c_out):,} bytes)")

print("\n--- 3. GENERATING HTML EXPORTS ---")
html_template = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
body {{
    font-family: Arial, sans-serif;
    line-height: 1.5;
    margin: 40px auto;
    max-width: 900px;
    padding: 0 20px;
    color: #111;
}}
table {{
    border-collapse: collapse;
    width: 100%;
    margin: 20px 0;
}}
th, td {{
    border: 1px solid #ddd;
    padding: 8px 12px;
    text-align: left;
}}
th {{
    background-color: #f2f2f2;
}}
code {{
    background-color: #f4f4f4;
    padding: 2px 4px;
    border-radius: 4px;
    font-family: Consolas, monospace;
}}
pre {{
    background-color: #f8f8f8;
    border: 1px solid #ddd;
    padding: 12px;
    overflow-x: auto;
}}
blockquote {{
    border-left: 4px solid #ccc;
    margin: 1.5em 10px;
    padding: 0.5em 10px;
    color: #555;
}}
</style>
</head>
<body>
{body}
</body>
</html>
"""

extensions = ['extra', 'tables', 'fenced_code', 'toc', 'sane_lists', 'nl2br']
md_files = [
    "BAGIAN_AWAL_PROPOSAL.md",
    "BAB_1_PENDAHULUAN.md",
    "BAB_2_LANDASAN_TEORI.md",
    "BAB_3_METODOLOGI_PENELITIAN.md",
    "MASTER_BIBLIOGRAPHY.md"
]

combined_html_parts = []
for fname in md_files:
    fpath = os.path.join(thesis_dir, fname)
    with open(fpath, 'r', encoding='utf-8') as f:
        content = f.read()
    html_body = markdown.markdown(content, extensions=extensions)
    out_name = os.path.splitext(fname)[0] + ".html"
    out_path = os.path.join(html_dir, out_name)
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(html_template.format(title=fname, body=html_body))
    print(f"Generated HTML: {out_path}")
    combined_html_parts.append(html_body)

full_html_out = os.path.join(html_dir, "FULL_PROPOSAL_MERGED.html")
with open(full_html_out, 'w', encoding='utf-8') as f:
    f.write(html_template.format(title="Proposal Lengkap AUV - Muh. Radhi Syafiq Ghanim. S", body="<hr style='margin:40px 0;'>".join(combined_html_parts)))
print(f"Generated merged HTML: {full_html_out}")

print("\n--- 4. SYNCHRONIZING WITH AUV_GitHub_Upload/Thesis ---")
for f in os.listdir(docx_dir):
    s = os.path.join(docx_dir, f)
    d = os.path.join(github_docx_dir, f)
    if os.path.isfile(s):
        shutil.copy2(s, d)

for f in os.listdir(html_dir):
    s = os.path.join(html_dir, f)
    d = os.path.join(github_html_dir, f)
    if os.path.isfile(s):
        shutil.copy2(s, d)

# Synchronize markdown source files
for f in os.listdir(thesis_dir):
    if f.endswith('.md'):
        s = os.path.join(thesis_dir, f)
        d = os.path.join(github_thesis_dir, f)
        shutil.copy2(s, d)

# Synchronize figures directory
figures_src = os.path.join(thesis_dir, "figures")
figures_html = os.path.join(html_dir, "figures")
figures_gh = os.path.join(github_thesis_dir, "figures")
figures_gh_html = os.path.join(github_html_dir, "figures")

for f_dir in [figures_html, figures_gh, figures_gh_html]:
    os.makedirs(f_dir, exist_ok=True)
    if os.path.exists(figures_src):
        for fig_f in os.listdir(figures_src):
            s_fig = os.path.join(figures_src, fig_f)
            d_fig = os.path.join(f_dir, fig_f)
            if os.path.isfile(s_fig):
                shutil.copy2(s_fig, d_fig)

shutil.copy2(logo_path, os.path.join(github_thesis_dir, "unhas_logo.png"))
print("All files and figures synchronized to AUV_GitHub_Upload/Thesis successfully!")
